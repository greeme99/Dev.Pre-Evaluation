"""Tests for XLSX and CSV findings export (FN-009, FR-013, TC-025, TC-026)."""

import csv
import io
from uuid import uuid4

import openpyxl
import pytest

from dev_pre_review.domain.enums import DrawingType, EvidenceKind, FindingStatus, Severity
from dev_pre_review.domain.models import (
    AnalysisJob,
    DrawingArtifact,
    Evidence,
    Finding,
    ReviewProject,
    utc_now,
)
from dev_pre_review.export.findings import FindingExporter
from dev_pre_review.storage.database import Database
from dev_pre_review.storage.repositories import (
    DrawingRepository,
    FindingRepository,
    JobRepository,
    ProjectRepository,
)


@pytest.fixture
def populated_db(tmp_path) -> tuple[Database, AnalysisJob, Finding]:
    db = Database.from_path(tmp_path / "test_export.db")
    db.upgrade()

    proj_repo = ProjectRepository(db.sessions)
    drawing_repo = DrawingRepository(db.sessions)
    job_repo = JobRepository(db.sessions)
    finding_repo = FindingRepository(db.sessions)

    p_id = uuid4()
    d_id = uuid4()
    j_id = uuid4()

    proj_repo.add(ReviewProject(id=p_id, name="내보내기 테스트 프로젝트", created_at=utc_now()))
    drawing_repo.add(
        DrawingArtifact(
            id=d_id,
            project_id=p_id,
            sha256="c" * 64,
            drawing_type=DrawingType.PDF,
            source_name="모터축_도면.pdf",
            parser_version="1.0.0",
            created_at=utc_now(),
        )
    )
    job = job_repo.add(
        AnalysisJob(
            id=j_id,
            project_id=p_id,
            drawing_id=d_id,
            created_at=utc_now(),
            updated_at=utc_now(),
        )
    )
    finding = Finding(
        id=uuid4(),
        job_id=j_id,
        status=FindingStatus.VIOLATION,
        severity=Severity.HIGH,
        confidence=0.95,
        finding_type="RULE-DIM-001",
        location="p.1 (10, 20)",
        message="외경 치수가 9.2mm로 기준 10.0mm에 미달합니다.",
        pixel_bbox=(10.0, 20.0, 30.0, 40.0),
        evidence=(
            Evidence(kind=EvidenceKind.RULE, ref_id="RULE-DIM-001"),
            Evidence(kind=EvidenceKind.CASE, ref_id="CASE-2025-01"),
        ),
        created_at=utc_now(),
    )
    finding_repo.add_many([finding])
    return db, job, finding


def test_export_findings_to_csv_with_utf8_bom(populated_db) -> None:
    db, job, finding = populated_db
    exporter = FindingExporter(db.sessions)

    csv_bytes = exporter.to_csv(job.id)

    # 1. UTF-8 BOM must be present at the start for Korean Excel compatibility
    assert csv_bytes.startswith(b"\xef\xbb\xbf")

    # 2. Decode and verify fixed columns
    text = csv_bytes.decode("utf-8-sig")
    reader = list(csv.reader(io.StringIO(text)))
    assert len(reader) >= 2  # Header + 1 finding row

    header = reader[0]
    expected_header = [
        "순번",
        "판정상태",
        "심각도",
        "규칙코드",
        "검토메시지",
        "도면위치",
        "근거목록",
        "도면상태",
    ]
    assert header == expected_header

    row = reader[1]
    assert row[1] == "위반"
    assert row[2] == "HIGH"
    assert row[3] == "RULE-DIM-001"
    assert "9.2mm" in row[4]
    assert "RULE:RULE-DIM-001" in row[6]
    assert row[7] == "정상"


def test_export_findings_to_xlsx_fixed_columns(populated_db) -> None:
    db, job, finding = populated_db
    exporter = FindingExporter(db.sessions)

    xlsx_bytes = exporter.to_xlsx(job.id)
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
    ws = wb.active
    assert ws is not None

    headers = [cell.value for cell in ws[1]]
    expected_headers = [
        "순번",
        "판정상태",
        "심각도",
        "규칙코드",
        "검토메시지",
        "도면위치",
        "근거목록",
        "도면상태",
    ]
    assert headers == expected_headers

    row_vals = [cell.value for cell in ws[2]]
    assert row_vals[1] == "위반"
    assert row_vals[3] == "RULE-DIM-001"
    assert row_vals[7] == "정상"


def test_export_findings_when_source_drawing_missing(populated_db) -> None:
    db, job, finding = populated_db
    # Pass a simulated drawing existence checker or check against missing artifact
    exporter = FindingExporter(db.sessions, source_exists_checker=lambda _d_id: False)

    csv_bytes = exporter.to_csv(job.id)
    text = csv_bytes.decode("utf-8-sig")
    reader = list(csv.reader(io.StringIO(text)))
    row = reader[1]
    # TC-026: 도면 부재 상태를 정직하게 출력
    assert row[7] == "원본 부재"
