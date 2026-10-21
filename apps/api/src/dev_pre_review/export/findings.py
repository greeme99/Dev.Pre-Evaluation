"""XLSX and UTF-8 BOM CSV exporter for review findings (FN-009, FR-013)."""

import csv
import io
from collections.abc import Callable
from uuid import UUID

import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session, sessionmaker

from dev_pre_review.domain.enums import EvidenceKind, FindingStatus
from dev_pre_review.domain.models import Finding
from dev_pre_review.storage.repositories import DrawingRepository, FindingRepository, JobRepository

HEADERS = [
    "순번",
    "판정상태",
    "심각도",
    "규칙코드",
    "검토메시지",
    "도면위치",
    "근거목록",
    "도면상태",
]

STATUS_KR_MAP = {
    FindingStatus.VIOLATION: "위반",
    FindingStatus.SUSPECTED: "의심",
    FindingStatus.INSUFFICIENT_INFORMATION: "정보 부족",
}
FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r", "\n")


def _escape_spreadsheet_cell(value: str) -> str:
    """Keep user-controlled strings inert when opened by spreadsheet software."""
    return f"'{value}" if value.startswith(FORMULA_PREFIXES) else value


class FindingExporter:
    """Exports findings for an analysis job to formatted XLSX or CSV."""

    def __init__(
        self,
        session_factory: sessionmaker[Session],
        source_exists_checker: Callable[[UUID], bool] | None = None,
    ) -> None:
        self.session_factory = session_factory
        self.source_exists_checker = source_exists_checker

    def _get_drawing_status(self, drawing_id: UUID) -> str:
        if self.source_exists_checker is not None:
            return "정상" if self.source_exists_checker(drawing_id) else "원본 부재"

        drawing_repo = DrawingRepository(self.session_factory)
        try:
            drawing_repo.get(drawing_id)
            return "정상"
        except Exception:
            return "원본 부재"

    def _format_row(self, idx: int, finding: Finding, drawing_status: str) -> list[str]:
        status_kr = STATUS_KR_MAP.get(finding.status, finding.status.value)
        severity_str = finding.severity.value.upper()

        bbox = finding.pixel_bbox or finding.cad_bbox
        if bbox:
            bbox_str = f"({bbox[0]:.1f}, {bbox[1]:.1f})~({bbox[2]:.1f}, {bbox[3]:.1f})"
        else:
            bbox_str = finding.location or "-"

        rule_code = next(
            (e.ref_id for e in finding.evidence if e.kind is EvidenceKind.RULE),
            finding.finding_type,
        )

        evidences_str = ", ".join(f"{e.kind.value.upper()}:{e.ref_id}" for e in finding.evidence)

        return [
            _escape_spreadsheet_cell(value)
            for value in [
                str(idx),
                status_kr,
                severity_str,
                rule_code,
                finding.message,
                bbox_str,
                evidences_str,
                drawing_status,
            ]
        ]

    def to_csv(self, job_id: UUID) -> bytes:
        job_repo = JobRepository(self.session_factory)
        finding_repo = FindingRepository(self.session_factory)

        job = job_repo.get(job_id)
        findings = finding_repo.list_for_job(job_id)
        drawing_status = self._get_drawing_status(job.drawing_id)

        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow(HEADERS)

        for idx, finding in enumerate(findings, start=1):
            writer.writerow(self._format_row(idx, finding, drawing_status))

        # UTF-8 with BOM for Korean Excel compatibility
        return output.getvalue().encode("utf-8-sig")

    def to_xlsx(self, job_id: UUID) -> bytes:
        job_repo = JobRepository(self.session_factory)
        finding_repo = FindingRepository(self.session_factory)

        job = job_repo.get(job_id)
        findings = finding_repo.list_for_job(job_id)
        drawing_status = self._get_drawing_status(job.drawing_id)

        wb = openpyxl.Workbook()
        ws = wb.active
        if ws is None:
            raise RuntimeError("Failed to create active worksheet")
        ws.title = "검토결과"

        ws.append(HEADERS)
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
        for cell in ws[1]:
            cell.font = header_font
            cell.fill = header_fill

        for idx, finding in enumerate(findings, start=1):
            ws.append(self._format_row(idx, finding, drawing_status))

        # Adjust column widths
        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_idx = col[0].column
            if isinstance(col_idx, int):
                col_letter = get_column_letter(col_idx)
                ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()
