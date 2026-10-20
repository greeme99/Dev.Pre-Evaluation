"""Excel reader extracting sheets, tabular structures, and embedded images with anchors."""

import hashlib
import zipfile
from pathlib import Path
from typing import Any
from uuid import uuid4

import openpyxl
from openpyxl.utils import get_column_letter

from dev_pre_review.domain.errors import DomainError
from dev_pre_review.knowledge.contracts import (
    ExtractedImage,
    ImportPreview,
    PreviewWarning,
    SheetPreview,
)


class ExcelReaderError(DomainError):
    code = "excel_reader_error"


class EncryptedExcelError(ExcelReaderError):
    code = "encrypted_excel_file"


class InvalidExcelError(ExcelReaderError):
    code = "invalid_excel_file"


class ExcelReader:
    def preview(self, path: Path) -> ImportPreview:
        if not path.exists():
            raise InvalidExcelError(f"File not found: {path}")

        # Check signature & encryption
        with open(path, "rb") as f:
            header = f.read(8)
            f.seek(0)
            file_hash = hashlib.sha256(f.read()).hexdigest()

        # OLE / Password encrypted signature check (d0 cf 11 e0)
        if header.startswith(b"\xd0\xcf\x11\xe0"):
            raise EncryptedExcelError(
                "Encrypted or legacy binary Excel (.xls) file is not supported"
            )

        if not zipfile.is_zipfile(path):
            raise InvalidExcelError("File is not a valid OpenXML spreadsheet (.xlsx)")

        try:
            wb = openpyxl.load_workbook(path, data_only=True)
        except Exception as exc:  # noqa: BLE001
            raise InvalidExcelError(f"Failed to open workbook: {exc}") from exc

        sheets: list[SheetPreview] = []
        images: list[ExtractedImage] = []
        warnings: list[PreviewWarning] = []

        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]

            # 1. Inspect sheet rows & columns
            rows_iter = ws.iter_rows(values_only=True)
            first_row = next(rows_iter, None)

            columns: list[str] = []
            if first_row:
                columns = [str(c).strip() for c in first_row if c is not None]

            sample_rows: list[dict[str, Any]] = []
            row_count = 0
            for r_vals in rows_iter:
                if not any(v is not None for v in r_vals):
                    continue
                row_count += 1
                if len(sample_rows) < 5 and columns:
                    row_dict = {
                        columns[idx]: r_vals[idx] for idx in range(min(len(columns), len(r_vals)))
                    }
                    sample_rows.append(row_dict)

            sheets.append(
                SheetPreview(
                    sheet_name=sheet_name,
                    columns=columns,
                    row_count=row_count,
                    sample_rows=sample_rows,
                )
            )

            # 2. Extract embedded images and anchor context
            ws_images = getattr(ws, "_images", [])
            for idx, img in enumerate(ws_images):
                anchor_cell = self._resolve_anchor_cell(img)
                context_text = self._build_context_text(ws, anchor_cell)

                img_data = img.ref.getvalue() if hasattr(img, "ref") else b""
                data_hash = hashlib.sha256(img_data).hexdigest() if img_data else ""

                images.append(
                    ExtractedImage(
                        image_id=f"img_{sheet_name}_{idx + 1}_{uuid4().hex[:6]}",
                        sheet_name=sheet_name,
                        anchor_cell=anchor_cell,
                        context_text=context_text,
                        data_hash=data_hash,
                    )
                )

        return ImportPreview(
            preview_id=uuid4(),
            file_sha256=file_hash,
            file_name=path.name,
            sheets=sheets,
            images=images,
            warnings=warnings,
        )

    def _resolve_anchor_cell(self, img: Any) -> str:
        anchor = getattr(img, "anchor", None)
        if isinstance(anchor, str):
            return anchor.upper()

        # Handle TwoCellAnchor or OneCellAnchor
        from_coord = getattr(anchor, "_from", None)
        if from_coord is not None:
            col = getattr(from_coord, "col", 0) + 1
            row = getattr(from_coord, "row", 0) + 1
            return f"{get_column_letter(col)}{row}"

        return "A1"

    def _build_context_text(self, ws: Any, anchor_cell: str) -> str:
        texts: list[str] = []

        # Check if anchor belongs to merged range
        for rng in ws.merged_cells.ranges:
            if anchor_cell in rng:
                top_left_cell = ws.cell(row=rng.min_row, column=rng.min_col)
                if top_left_cell.value is not None:
                    texts.append(str(top_left_cell.value))
                break

        # Coordinate based search around anchor (adjacent cells)
        try:
            cell = ws[anchor_cell]
            if cell.value is not None and str(cell.value) not in texts:
                texts.append(str(cell.value))

            r, c = cell.row, cell.column
            for dr in (-1, 0, 1):
                for dc in (-1, 0, 1):
                    if dr == 0 and dc == 0:
                        continue
                    adj_r, adj_c = r + dr, c + dc
                    if adj_r >= 1 and adj_c >= 1:
                        adj_val = ws.cell(row=adj_r, column=adj_c).value
                        if adj_val is not None:
                            val_str = str(adj_val).strip()
                            if val_str and val_str not in texts:
                                texts.append(val_str)
        except Exception:  # noqa: BLE001, S110
            pass

        return " | ".join(texts)
